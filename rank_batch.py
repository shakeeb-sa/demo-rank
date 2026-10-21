from patchright.sync_api import sync_playwright
import time, requests, glob, os, sys
from datetime import datetime
from openpyxl import Workbook, load_workbook

DOMAIN = "virtueam"
MAX_MATCHES = 5
HEADER = ["keyword", "found", "total_scanned", "matches_found",
          "rank1", "url1", "rank2", "url2", "rank3", "url3",
          "rank4", "url4", "rank5", "url5"]


def resolve(link):
    try:
        r = requests.get(link, allow_redirects=False, timeout=8)
        return r.headers.get("Location", link)
    except Exception as e:
        return f"ERR:{e}"


def read_keywords(path):
    if path.lower().endswith(".txt"):
        with open(path, "r", encoding="utf-8") as f:
            return [l.strip() for l in f if l.strip()]
    if path.lower().endswith(".xlsx"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        return [str(r[0]).strip() for r in rows[1:]
                if r and r[0] and str(r[0]).strip()]
    raise ValueError(f"Unsupported input file: {path}")


def find_latest():
    files = glob.glob("results_*.xlsx")
    return max(files, key=os.path.getmtime) if files else None


def load_completed(path):
    if not path:
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    done = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            done[str(row[0]).strip()] = row
    wb.close()
    return done


def process(page, kw):
    page.goto("https://www.google.com", wait_until="domcontentloaded")
    box = page.locator("textarea[name=q], input[name=q]").first
    box.click()
    box.fill(kw)
    box.press("Enter")
    try:
        page.wait_for_selector("#rso", timeout=15000)
    except Exception:
        if "/sorry/" in page.url:
            print("  CAPTCHA — solve it in the window, then wait…")
            page.wait_for_url("**/search**", timeout=180000)
            page.wait_for_selector("#rso", timeout=15000)
        else:
            raise
    time.sleep(2)
    page.keyboard.press("Alt+Shift+A")
    time.sleep(3)
    last = 0
    for _ in range(25):
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        time.sleep(1)
        h = page.evaluate("document.body.scrollHeight")
        if h == last:
            time.sleep(2)
            page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            if page.evaluate("document.body.scrollHeight") == h:
                break
        last = h
    results = page.eval_on_selector_all("#rso a:has(h3)", """
        els => els.map(a => ({
            href: a.href,
            cite: a.querySelector('cite') ? a.querySelector('cite').innerText : ''
        }))
    """)
    total = len(results)
    if total == 0:
        return ("ERROR", 0, [])
    hits = [(i, r['href']) for i, r in enumerate(results, 1)
            if DOMAIN in r['cite'].lower()]
    if not hits:
        return ("No", total, [])
    resolved = [(pos, resolve(href)) for pos, href in hits[:MAX_MATCHES]]
    return ("Yes", total, resolved)


def build_row(kw, found, total, resolved):
    row = [kw, found, total, len(resolved)]
    for i in range(MAX_MATCHES):
        if i < len(resolved):
            row.extend([resolved[i][0], resolved[i][1]])
        else:
            row.extend([None, None])
    return row


def main():
    if len(sys.argv) != 2:
        print("Usage: python rank_batch.py keywords.xlsx|keywords.txt")
        sys.exit(1)
    input_file = sys.argv[1]
    keywords = read_keywords(input_file)
    print(f"Loaded {len(keywords)} keywords from {input_file}")

    latest = find_latest()
    completed = load_completed(latest) if latest else {}
    retry = {k for k, r in completed.items() if r[1] == "ERROR"}
    if completed:
        print(f"Resumed from {latest}: "
              f"{len(completed) - len(retry)} done, {len(retry)} to retry")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_file = f"results_{ts}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for kw, row in completed.items():
        if row[1] in ("Yes", "No"):
            ws.append(list(row))
    wb.save(out_file)
    print(f"Output: {out_file}")

    pending = [k for k in keywords
               if k not in completed or k in retry]
    print(f"Pending: {len(pending)} keywords")
    if not pending:
        print("Nothing to do.")
        return

    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            "./chrome-profile", headless=False, no_viewport=True)
        page = ctx.pages[0]
        for idx, kw in enumerate(pending, 1):
            print(f"[{idx}/{len(pending)}] {kw}")
            try:
                found, total, resolved = process(page, kw)
            except Exception as e:
                print(f"  ERROR: {e}")
                found, total, resolved = ("ERROR", 0, [])
            ws.append(build_row(kw, found, total, resolved))
            wb.save(out_file)
            print(f"  {found} | scanned={total} | hits={len(resolved)}")
            time.sleep(2)
        ctx.close()
    wb.save(out_file)
    print(f"DONE — saved {out_file}")


if __name__ == "__main__":
    main()